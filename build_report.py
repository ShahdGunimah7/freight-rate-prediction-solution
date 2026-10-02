"""Builds reports/Freight_Rate_Report.docx from reports/metrics.json and the figures.
Run after run_pipeline.py, the figures step and score.py."""
from __future__ import annotations

import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from src.freight_rate import config

M = json.load(open(config.REPORTS_DIR / "metrics.json"))
FIG = config.FIGURES_DIR
CHART = config.ROOT / "scorer_results" / "candidate_december.png"
ACCENT = RGBColor(0x06, 0x4A, 0x56)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
for s in ("Heading 1", "Heading 2"):
    doc.styles[s].font.color.rgb = ACCENT
for section in doc.sections:
    section.left_margin = section.right_margin = Inches(0.9)
    section.top_margin = section.bottom_margin = Inches(0.8)


def para(text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)


def figure(path, caption, width=6.4):
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True
    c.runs[0].font.size = Pt(9)


def table(header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()


dq = M["data_quality"]
folds = M["backtests"]
final = "LightGBM + additive calendar effects (final)"
fA, fB = (folds[k]["models"][final]["clean_labels"] for k in folds)

doc.add_heading("Freight Rate Prediction — Machine Learning Engineer Assessment", 0)
para("Author: Shahd Ghunimah")

doc.add_heading("1. Summary", 1)
para(
    f"A LightGBM model trained on log(posted_rate) predicts the 12,000 Nov–Dec 2025 validation loads. "
    f"Validation uses time-based backtests that mirror the real task (train on everything up to the first month of a quarter, "
    f"forecast the next two months). The final model reaches {fA['MAPE_%']:.2f}% and {fB['MAPE_%']:.2f}% MAPE "
    f"(MAE ${fA['MAE']:.0f} / ${fB['MAE']:.0f}) on clean labels in the two backtest folds, against 8–9% for a rate-per-mile baseline "
    f"and 3–4.5% for a log-linear regression."
)

doc.add_heading("2. Key findings from the data", 1)
bullet(" Distance explains most of the rate (rate ∝ distance^0.87). Reefer is about 13% and Flatbed about 8% more expensive than Dry Van, and heavier loads cost more.", "Structure.")
bullet(" Rates ramp up 5–6% in the last ~5 weeks of every quarter, then reset on the 1st of the next quarter. Each quarter also sits about 1.5–2% above the previous one. Nov–Dec is the end of Q4, so the forecast has to capture this ramp.", "Quarter-end seasonality.")
bullet(" market_index is a daily market signal with a strong weekly cycle (peaks Wed/Thu). It is lower in Nov–Dec (~0.93) than the training average (1.08), but still inside the training range.", "Market index.")
bullet(f" {len(dq['validation_cities_not_in_training'])} validation cities never appear in training ({', '.join(dq['validation_cities_not_in_training'])}), about 12% of validation loads. The model therefore uses coordinates, not city IDs.", "Unseen cities.")
figure(FIG / "03_quarter_end_ramp.png", "Figure 1 – Daily mean residual of a simple log-linear model: rates ramp into each quarter end.")
figure(FIG / "04_market_index.png", "Figure 2 – market_index over time and its weekly cycle.")

doc.add_heading("3. Data-quality issues and how they were handled", 1)
table(
    ["Issue", "Train", "Validation", "Handling"],
    [
        ["Corrupted labels (rate ×0.17–0.45 or ×2.2–4.6)", f"{dq['label_outliers_removed']} (1.4%)", "n/a", "Flagged by an iteratively refit robust log-linear model (|log residual| > 0.4) and excluded from training"],
        ["Negative weight", dq["negative_weight"], dq["validation_negative_weight"], "Sign error, so absolute value taken"],
        ["Missing weight", dq["missing_weight"], dq["validation_missing_weight"], "Median weight of the same equipment type (training statistics)"],
        ["Missing market_index", dq["missing_market_index"], dq["validation_missing_market_index"], "Mean market_index of the same date (within-day std 0.025 vs 0.17 across days)"],
        ["Weight capped at 47,500 lb", dq["weight_at_47500_cap"], dq["validation_weight_at_47500_cap"], "Kept: it is a plausible legal/load cap and tree models handle it"],
        ["Synthetic-looking coordinates", "—", "—", "Coordinates are consistent per city, so they are used as-is. Distance is consistent per lane and is used as the main feature"],
    ],
)
figure(FIG / "02_label_outliers.png", "Figure 3 – Actual / expected rate. Corrupted labels form two separate clusters, far from the clean peak. Dashed lines = removal threshold.")

doc.add_heading("4. Train / test split and validation approach", 1)
para(
    "The labelled data covers Jan–Oct 2025 and the loads to predict are Nov–Dec 2025, so this is a forecasting problem. "
    "A random split would leak the market conditions of every day into training and give an over-optimistic score. "
    "So I used forward-in-time backtests that copy the real setup:"
)
for name, f in folds.items():
    bullet(f" train {f['train_window']} ({f['n_train']:,} clean loads), test {f['test_window']} ({f['n_test']:,} loads).", name + ":")
para(
    "In both folds the model sees only the first month of the target quarter, exactly like Oct → Nov–Dec. "
    "Cleaning statistics and outlier flags are re-fit inside each fold's training window, so no information from the test window leaks in. "
    "Metrics are reported on clean labels (the real target) and on all labels. "
    "Hyper-parameters (tree size, number of rounds) were chosen on these folds. "
    "The final model is then retrained on all clean Jan–Oct data."
)
para("Two extra checks:", None)
r = M["random_split_80_20_for_comparison"]
bullet(f" a random 80/20 split gives {r['MAPE_%']:.2f}% MAPE, lower than the honest time-based estimate. This confirms that a random split is optimistic.", "Random split (for comparison only):")
u = M["unseen_city_check"]
bullet(
    f" 8 cities were removed entirely from Fold B's training data. Loads touching them scored {u['rows_with_unseen_city']['MAPE_%']:.2f}% MAPE vs "
    f"{u['rows_with_seen_cities']['MAPE_%']:.2f}% for the other loads, so the model generalises to new cities.",
    "Unseen-city test:",
)

doc.add_heading("5. Model", 1)
para(
    "LightGBM gradient boosting on log(posted_rate), averaged over 3 seeds. Features: log distance, equipment, weight, market_index, "
    "quote_signal, pickup/delivery coordinates, bearing, and calendar features (quarter, days to quarter end, day of week, day of month). "
    "The log target turns the multiplicative price structure into an additive one and optimises relative error."
)
para(
    "The key design choice is LightGBM interaction constraints: each calendar feature is kept in its own group, so its effect is additive in "
    "log space (a multiplicative uplift on the rate). The quarter-end ramp learned in Q1–Q3 then transfers cleanly to the end of Q4. "
    "Without the constraint, trees mix the quarter with days-to-quarter-end and cannot extrapolate the ramp (2.62% vs 1.78% MAPE in Fold A). "
    "Small trees (4 leaves, 3,000 rounds at learning rate 0.03) generalised better than deeper ones."
)

doc.add_heading("6. Results", 1)
rows = []
for name in folds["Fold A"]["models"]:
    a = folds["Fold A"]["models"][name]["clean_labels"]
    b = folds["Fold B"]["models"][name]["clean_labels"]
    rows.append([name, f"{a['MAPE_%']:.2f}%", f"${a['MAE']:.1f}", f"{b['MAPE_%']:.2f}%", f"${b['MAE']:.1f}", f"{b['R2']:.4f}"])
table(["Model", "Fold A MAPE", "Fold A MAE", "Fold B MAPE", "Fold B MAE", "Fold B R²"], rows)
aA = folds["Fold A"]["models"][final]["all_labels"]
aB = folds["Fold B"]["models"][final]["all_labels"]
para(
    f"Including the corrupted labels in the test set, the final model's MAE is ${aA['MAE']:.0f} / ${aB['MAE']:.0f}. "
    "The difference comes entirely from labels that are off by a factor of 2–5, which no model can or should predict."
)
figure(FIG / "05_backtest_mape.png", "Figure 4 – Backtest MAPE on clean labels.")

doc.add_heading("7. December prediction chart (score.py)", 1)
para(
    "Fixed inputs: Lexington → Fort Wayne, 360 mi, Dry Van, 32,000 lb. Only the date changes. The chart input file has no market "
    "signals, so for each day I used the daily mean market_index and median quote_signal observed in validation.csv for that date. "
    "The curve rises about 6% through December (the Q4 quarter-end ramp), with a weekly ripple from the market_index cycle."
)
figure(CHART, "Figure 5 – Output of the provided score.py.", width=6.6)

doc.add_heading("8. Limitations and next steps", 1)
bullet(" The training data contains no Thanksgiving or Christmas periods, so holiday effects cannot be learned.", "Holidays:")
bullet(" The quarter-end ramp is assumed to behave in Q4 as in Q1–Q3. This matched well in both backtests.", "Seasonality assumption:")
bullet(" Add prediction intervals (quantile LightGBM) and monitor the corrupted-label rate in production.", "Next:")

out = config.REPORTS_DIR / "Freight_Rate_Report.docx"
doc.save(out)
print(f"Report written to {out}")
